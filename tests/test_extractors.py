from email.message import EmailMessage
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from efds.ingestion.extractors import (
    extract_docx,
    extract_eml,
    extract_txt,
    extract_xlsx,
    infer_document_type,
)


def test_txt_and_markdown_extraction_handles_bom(tmp_path: Path) -> None:
    txt = tmp_path / "notes.txt"
    markdown = tmp_path / "readme.md"
    txt.write_bytes("\ufeffHello text".encode("utf-8"))
    markdown.write_text("# Heading\n\nBody", encoding="utf-8")
    assert extract_txt(txt).text == "Hello text"
    assert extract_txt(markdown).text == "# Heading\n\nBody"
    assert infer_document_type(markdown) == "markdown"


def test_docx_extraction_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "minutes.docx"
    document = Document()
    document.add_paragraph("Opening paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Action"
    document.add_paragraph("Closing paragraph")
    document.save(path)
    text = extract_docx(path).text
    assert "Opening paragraph" in text
    assert "Owner | Action" in text
    assert text.index("Opening paragraph") < text.index("Owner | Action") < text.index("Closing paragraph")


def test_eml_extraction_includes_headers_body_and_attachment_names(tmp_path: Path) -> None:
    path = tmp_path / "message.eml"
    message = EmailMessage()
    message["Subject"] = "Welcome"
    message["From"] = "Chair <chair@example.com>"
    message["To"] = "team@example.com"
    message["Cc"] = "finance@example.com"
    message["Date"] = "Tue, 11 Aug 2026 10:00:00 +0100"
    message.set_content("Plain body")
    message.add_attachment(b"ignored", maintype="application", subtype="octet-stream", filename="file.bin")
    path.write_bytes(bytes(message))
    result = extract_eml(path)
    assert "Subject: Welcome" in result.text
    assert "Plain body" in result.text
    assert result.metadata["sender"] == "Chair <chair@example.com>"
    assert result.metadata["recipients"] == ["team@example.com"]
    assert result.metadata["attachment_names"] == ["file.bin"]


def test_xlsx_extraction_includes_sheet_names_and_non_empty_cells(tmp_path: Path) -> None:
    path = tmp_path / "officers.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Officers"
    worksheet.append(["Name", "Role"])
    worksheet.append(["Siheon Lee", "Chair"])
    workbook.save(path)
    result = extract_xlsx(path)
    assert "--- Sheet: Officers ---" in result.text
    assert "Siheon Lee\tChair" in result.text

